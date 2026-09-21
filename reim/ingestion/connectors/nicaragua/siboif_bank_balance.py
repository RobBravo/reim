"""SIBOIF — Nicaragua banking system balance sheet.

Three system-wide totals (assets, liabilities, equity), monthly, read from
the SISTEMA_BANCARIO sheet of SIBOIF's published balance sheet workbook.
Verified to satisfy Activo = Pasivo + Patrimonio for every period this
connector produces.

See docs/sources.md's SIBOIF entry and
docs/superpowers/specs/2026-09-21-siboif-bank-balance-design.md for the
research this is built from.
"""

from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal
from typing import Any

import openpyxl

from reim.core.constants import CheckSeverity, CheckType, Frequency
from reim.core.exceptions import TransformationError
from reim.domain.observations.periods import parse_period
from reim.domain.pipelines.models import NormalizedObservation, QualityResult, RawDataset
from reim.ingestion.base import BaseConnector
from reim.ingestion.http import ensure_ok, fetch, http_client

SHEET_NAME = "SISTEMA_BANCARIO"

#: Column-A text -> REIM indicator code. Order doesn't matter; this is the
#: authority on which three of the sheet's 95 rows this connector reads.
ROW_INDICATORS: dict[str, str] = {
    "Activo": "ni_bank_system_total_assets_monthly",
    "Pasivo": "ni_bank_system_total_liabilities_monthly",
    "PATRIMONIO": "ni_bank_system_total_equity_monthly",
}

#: Tolerance for the accounting identity, in the source's own thousands-of-NIO
#: scale (matching its own four-decimal precision) -- not the whole-córdoba
#: scale the connector stores. Measured against the real fixture on
#: 2026-09-20: 20 of 92 published months carry a rounding artifact of up to
#: 0.0002 in this scale (SIBOIF's own published figures, not this
#: connector's arithmetic), so the tolerance must clear that, not just the
#: precision the source claims.
IDENTITY_TOLERANCE = Decimal("0.001")

#: The scale the sheet declares for itself, above its own header row (cell A8
#: in the recording measured on 2026-09-21: "(Expresado en miles de
#: Córdobas)"). The x1,000 conversion to whole córdobas is only valid while
#: that stays true, and nothing downstream would notice if it stopped: the
#: balance identity holds under any uniform rescaling, ``max_value`` is null
#: and no ``max_period_change_pct`` is declared. This literal is the guard.
EXPECTED_UNIT_NOTE = "miles de Córdobas"


class SIBOIFBankBalanceConnector(BaseConnector):
    """System-wide balance-sheet totals for Nicaragua's banking sector."""

    connector_key = "siboif_bank_balance"
    version = "1.0.0"
    expected_frequency = Frequency.MONTHLY

    async def extract(self) -> RawDataset:
        """Fetch the static balance-sheet workbook.

        One request. The file is refreshed in place by SIBOIF (no date or
        version in the URL), so each run re-reads the full history.

        Raises:
            ExtractionError: The file was unreachable or not the expected
                content type.
        """
        url = (
            f"{str(self.source.base_url).rstrip('/')}/sites/default/files/documentos/"
            "serie-informes-excel/bancos/ib_balance_general_0.xlsx"
        )
        retrieved_at = self.now()

        async with http_client(user_agent=self.source.user_agent) as client:
            response = await fetch(client, url)
            ensure_ok(
                response,
                expected_content_type=(
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                ),
            )
            content = response.content

        return RawDataset(
            source_key=self.connector_key,
            retrieved_at=retrieved_at,
            source_url=url,
            payload=content,
            content_type=response.headers.get("content-type"),
            http_status=response.status_code,
            metadata={"sheet": SHEET_NAME},
        )

    def transform(self, raw: RawDataset) -> list[NormalizedObservation]:
        """Normalize the SISTEMA_BANCARIO sheet into three monthly series.

        Pure function of ``raw``.

        Raises:
            TransformationError: The sheet, header row or one of the three
                expected rows could not be found, the sheet no longer declares
                its figures in thousands of córdobas, or no observations
                resulted.
        """
        workbook = openpyxl.load_workbook(io.BytesIO(raw.payload), read_only=True, data_only=True)
        try:
            sheet = workbook[SHEET_NAME]
        except KeyError as exc:
            msg = f"Sheet {SHEET_NAME!r} not found in workbook"
            raise TransformationError(msg) from exc

        header_row = _find_header_row(sheet)
        self._assert_declared_unit(sheet, header_row)
        month_columns = _month_columns(sheet, header_row)

        observations: list[NormalizedObservation] = []
        for label, indicator_code in ROW_INDICATORS.items():
            row_cells = _find_row_by_label(sheet, label)
            for column_index, date_str in month_columns:
                value = row_cells[column_index]
                if value is None:
                    continue
                period = parse_period(_to_monthly_label(date_str), Frequency.MONTHLY)
                thousands_value = Decimal(str(value))
                observations.append(
                    NormalizedObservation(
                        country_iso3="NIC",
                        indicator_code=indicator_code,
                        source_key=self.connector_key,
                        period=period,
                        unit="current NIO",
                        retrieved_at=raw.retrieved_at,
                        source_url=raw.source_url,
                        value_numeric=thousands_value * 1000,
                        currency_code="NIO",
                        raw_metadata={"siboif_value_thousands_nio": str(thousands_value)},
                    )
                )

        if not observations:
            msg = "SIBOIF connector produced zero observations"
            raise TransformationError(msg)

        return observations

    def _assert_declared_unit(self, sheet: Any, header_row: int) -> None:
        """Fail loudly unless the sheet still declares its own figures in thousands.

        Guards the one silent-corruption mode ``validate()`` structurally
        cannot see: the accounting identity holds under any uniform rescaling,
        so a republication in whole córdobas or in millions would make every
        stored value wrong by 1,000x without failing a single check.
        """
        for row in sheet.iter_rows(min_row=1, max_row=header_row, max_col=1, values_only=True):
            if isinstance(row[0], str) and EXPECTED_UNIT_NOTE in row[0]:
                return
        msg = (
            f"Sheet {SHEET_NAME} does not declare {EXPECTED_UNIT_NOTE!r} above its "
            f"header row; the x1000 conversion to whole córdobas is only valid "
            f"while the source publishes thousands of córdobas."
        )
        raise TransformationError(msg, source_key=self.source.key)

    def validate(self, observations: list[NormalizedObservation]) -> list[QualityResult]:
        """Assert Activo == Pasivo + Patrimonio for every period produced."""
        by_period: dict[str, dict[str, Decimal]] = {}
        for obs in observations:
            if obs.value_numeric is None:
                continue
            by_period.setdefault(obs.period.label, {})[obs.indicator_code] = obs.value_numeric

        broken = []
        for period, values in sorted(by_period.items()):
            assets = values.get("ni_bank_system_total_assets_monthly")
            liabilities = values.get("ni_bank_system_total_liabilities_monthly")
            equity = values.get("ni_bank_system_total_equity_monthly")
            if assets is None or liabilities is None or equity is None:
                continue
            if abs(assets - liabilities - equity) > IDENTITY_TOLERANCE * 1000:
                broken.append(period)

        if not broken:
            return [
                QualityResult.passed(
                    "siboif_balance_identity",
                    CheckType.CONSISTENCY,
                    f"Activo = Pasivo + Patrimonio holds for all {len(by_period)} period(s)",
                    expected_value="0 beyond tolerance",
                    actual_value="0",
                )
            ]
        return [
            QualityResult.failure(
                "siboif_balance_identity",
                CheckType.CONSISTENCY,
                CheckSeverity.ERROR,
                f"{len(broken)} period(s) break the balance identity: {', '.join(broken[:5])}",
                expected_value="0 beyond tolerance",
                actual_value=str(len(broken)),
            )
        ]


def _find_header_row(sheet: Any) -> int:
    """Return the row index whose first cell is 'Descripción'."""
    for row in sheet.iter_rows(min_row=1, max_row=20):
        if row[0].value == "Descripción":
            return int(row[0].row)
    msg = "Could not find the 'Descripción' header row"
    raise TransformationError(msg)


def _month_columns(sheet: Any, header_row: int) -> list[tuple[int, str]]:
    """Return (0-based column index, date string) for every dated column."""
    columns = []
    for cell in sheet[header_row]:
        if cell.column == 1:
            continue
        if isinstance(cell.value, str) and "/" in cell.value:
            columns.append((cell.column - 1, cell.value))
    return columns


def _find_row_by_label(sheet: Any, label: str) -> list[Any]:
    """Return the full row of values whose column-A text exactly matches ``label``."""
    for row in sheet.iter_rows(values_only=True):
        if row and row[0] == label:
            return list(row)
    msg = f"Could not find a row labelled {label!r} in {SHEET_NAME}"
    raise TransformationError(msg)


def _to_monthly_label(date_str: str) -> str:
    """Convert a day-first 'DD/MM/YYYY' string to REIM's 'YYYY-MM' monthly label."""
    parsed = datetime.strptime(date_str, "%d/%m/%Y")
    return f"{parsed.year:04d}-{parsed.month:02d}"
